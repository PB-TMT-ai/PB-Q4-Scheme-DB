"""Integration tests for data loading.

These tests require a valid Excel file in data/ to run.
Skip gracefully if no data file is present.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

DATA_DIR = Path(__file__).resolve().parent.parent.parent / "data"


@pytest.fixture
def has_data_file() -> bool:
    """Check if a data file exists for integration tests."""
    files = list(DATA_DIR.glob("*.xlsx"))
    if not files:
        pytest.skip("No Excel file in data/ — skipping integration tests")
    return True


class TestDataLoading:
    """Integration tests for data loading pipeline."""

    def test_report_script_imports(self) -> None:
        """Verify the report script can be imported without errors."""
        from scripts.generate_slab_report import (
            SLAB_CONFIG, assign_slab, format_indian,
            get_next_slab, points_to_next,
        )
        assert len(SLAB_CONFIG) > 0
        assert assign_slab(0) == "Unqualified"
        assert get_next_slab("Unqualified") == "Slab A"
        assert points_to_next(500, "Unqualified") == 250.0
        assert get_next_slab("Slab E") is None

    def test_load_data_with_file(self, has_data_file: bool) -> None:
        """Test data loading returns a valid DataFrame."""
        from scripts.generate_slab_report import load_data
        df = load_data()
        assert len(df) > 0
        assert "Qualified Slab" in df.columns
        assert "Qualified Points" in df.columns

    def test_self_counter_excluded(self, has_data_file: bool) -> None:
        """Verify self-counter dealers are excluded from loaded data."""
        from scripts.generate_slab_report import load_data
        df = load_data()
        if "Self Counter" in df.columns:
            assert (df["Self Counter"] != "Yes").all()

    def test_column_mapping_applied(self, has_data_file: bool) -> None:
        """Verify Excel columns are renamed to standard names."""
        from scripts.generate_slab_report import load_data
        df = load_data()
        expected_cols = ["Dealer Name", "State", "District", "Zone",
                         "Qualified Volume", "Qualified Points"]
        for col in expected_cols:
            assert col in df.columns, f"Expected column '{col}' not found"

    def test_next_slab_columns_exist(self, has_data_file: bool) -> None:
        """Verify Next Upgrade Slab and Points to Next Slab columns are computed."""
        from scripts.generate_slab_report import load_data
        df = load_data()
        assert "Next Upgrade Slab" in df.columns
        assert "Points to Next Slab" in df.columns

    def test_slab_e_has_no_next(self, has_data_file: bool) -> None:
        """Verify Slab E dealers have None for Next Upgrade Slab."""
        from scripts.generate_slab_report import load_data
        df = load_data()
        slab_e = df[df["Qualified Slab"] == "Slab E"]
        if len(slab_e) > 0:
            assert slab_e["Next Upgrade Slab"].isna().all()

    def test_generate_report(self, has_data_file: bool) -> None:
        """Test report generation produces an output file."""
        from scripts.generate_slab_report import generate_report, load_data
        df = load_data()
        output_path = generate_report(df)
        assert output_path.exists()
        assert output_path.suffix == ".xlsx"

    def test_report_has_all_sheets(self, has_data_file: bool) -> None:
        """Verify report contains all four expected sheets."""
        from openpyxl import load_workbook
        from scripts.generate_slab_report import generate_report, load_data
        df = load_data()
        output_path = generate_report(df)
        wb = load_workbook(str(output_path))
        expected = {"Slab Summary", "Dealer Detail", "Near-Upgrade", "Distributor Performance"}
        assert expected == set(wb.sheetnames)
