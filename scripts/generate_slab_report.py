"""
Standalone Slab Report Generator
================================
Generates a formatted Excel report with slab analysis based on
qualified points from the source data. Outputs to .workspace/ directory.

Usage:
    python scripts/generate_slab_report.py
"""

from __future__ import annotations

import glob
import os
import sys
from pathlib import Path
from typing import Optional

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Add project root to path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from src.lib.logger import error as log_error
from src.lib.logger import info as log_info

# ---------------------------------------------------------------------------
# CONSTANTS (mirrored from app.py)
# ---------------------------------------------------------------------------

DATA_DIR: str = str(PROJECT_ROOT / "data")
OUTPUT_DIR: str = str(PROJECT_ROOT / ".workspace")
SHEET_NAME: str = "Sheet1"
HEADER_ROW: int = 0

# Column name mapping: Excel column → internal standard name
COLUMN_MAP: dict[str, str] = {
    "Retailer Name": "Dealer Name",
    "State Name": "State",
    "District Name": "District",
    "Shop vol.": "Shop Volume",
    "Site vol.": "Site Volume",
    "Total vol. under scheme": "Qualified Volume",
    "Total site vol.": "Total Site Volume",
    "Points": "Qualified Points",
    "Current gift": "Gift",
    "Current gift slab": "Current Slab",
    "Distributor self-counter (Yes/No)": "Self Counter",
    "Jan + Feb+Mar": "Q4 Volume",
    "# Unique Site >200 MT": "Unique Site >200 MT",
}

SLAB_CONFIG: list[dict] = [
    {"slab": "Unqualified", "slab_code": "-", "range": "0 – 749",
     "lower": 0, "upper": 750, "gift": "No Gift",
     "gift_full": "No Gift", "category": "Unqualified",
     "volume_mt": 0, "color": "#94a3b8", "color_light": "#f1f5f9"},
    {"slab": "Slab A", "slab_code": "A", "range": "750 – 2,999",
     "lower": 750, "upper": 3000, "gift": "Foot Massager",
     "gift_full": "Foot massager", "category": "A",
     "volume_mt": 30, "color": "#f59e0b", "color_light": "#fef3c7"},
    {"slab": "Slab B", "slab_code": "B", "range": "3,000 – 4,199",
     "lower": 3000, "upper": 4200, "gift": "Sony Sound Bar",
     "gift_full": "Sony - Sound bar, woofer and speakers",
     "category": "B", "volume_mt": 120,
     "color": "#6366f1", "color_light": "#e0e7ff"},
    {"slab": "Slab C", "slab_code": "C", "range": "4,200 – 6,799",
     "lower": 4200, "upper": 6800, "gift": "Robot Vacuum",
     "gift_full": "Robot Vacuum cleaner", "category": "C",
     "volume_mt": 168, "color": "#10b981", "color_light": "#d1fae5"},
    {"slab": "Slab D", "slab_code": "D", "range": "6,800 – 7,499",
     "lower": 6800, "upper": 7500, "gift": "Apple iPad",
     "gift_full": "Apple iPad", "category": "D",
     "volume_mt": 272, "color": "#3b82f6", "color_light": "#dbeafe"},
    {"slab": "Slab E", "slab_code": "E", "range": "7,500+",
     "lower": 7500, "upper": float("inf"), "gift": "Washing Machine",
     "gift_full": "Samsung front-load washing machine",
     "category": "E", "volume_mt": 300,
     "color": "#ec4899", "color_light": "#fce7f3"},
]

SLAB_ORDER: list[str] = [s["slab"] for s in SLAB_CONFIG]
SLAB_COLORS_LIGHT: dict[str, str] = {s["slab"]: s["color_light"] for s in SLAB_CONFIG}
SLAB_CODE_MAP: dict[str, str] = {s["slab_code"]: s["slab"] for s in SLAB_CONFIG}

# Next-slab mapping (mirrored from app.py)
NEXT_SLAB_MAP: dict[str, Optional[str]] = {}
NEXT_SLAB_THRESHOLD: dict[str, Optional[float]] = {}
NEXT_SLAB_VOLUME: dict[str, Optional[float]] = {}
for _i, _cfg in enumerate(SLAB_CONFIG):
    NEXT_SLAB_MAP[_cfg["slab"]] = (
        SLAB_CONFIG[_i + 1]["slab"] if _i + 1 < len(SLAB_CONFIG) else None
    )
    NEXT_SLAB_THRESHOLD[_cfg["slab"]] = (
        SLAB_CONFIG[_i + 1]["lower"] if _i + 1 < len(SLAB_CONFIG) else None
    )
    NEXT_SLAB_VOLUME[_cfg["slab"]] = (
        SLAB_CONFIG[_i + 1]["volume_mt"] if _i + 1 < len(SLAB_CONFIG) else None
    )

# ---------------------------------------------------------------------------
# HELPERS (mirrored from app.py)
# ---------------------------------------------------------------------------


def assign_slab(points: float) -> str:
    """Assign a slab label based on qualified points.

    Args:
        points: Qualified points value.

    Returns:
        Slab label string.
    """
    if pd.isna(points):
        points = 0.0
    points = float(points)
    for cfg in SLAB_CONFIG:
        if cfg["lower"] <= points < cfg["upper"]:
            return cfg["slab"]
    return SLAB_CONFIG[0]["slab"]


def get_next_slab(current: str) -> Optional[str]:
    """Look up the next higher slab from NEXT_SLAB_MAP.

    Args:
        current: Current slab label.

    Returns:
        Next slab label, or None if already at maximum.
    """
    return NEXT_SLAB_MAP.get(current)


def points_to_next(points: float, current: str) -> Optional[float]:
    """Calculate the points gap to reach the next slab tier.

    Args:
        points: Current qualified points.
        current: Current slab label.

    Returns:
        Points needed to reach next tier, or None if at max slab.
    """
    threshold = NEXT_SLAB_THRESHOLD.get(current)
    if threshold is None:
        return None
    gap = threshold - float(points)
    return max(gap, 0.0)


def format_indian(number: float, prefix: str = "", decimal: int = 0) -> str:
    """Format a number with Indian comma grouping.

    Args:
        number: The numeric value to format.
        prefix: Optional prefix string.
        decimal: Number of decimal places.

    Returns:
        Formatted string with Indian-style commas.
    """
    if pd.isna(number):
        return f"{prefix}0"
    number = round(float(number), decimal)
    is_negative = number < 0
    number = abs(number)

    if decimal > 0:
        int_part, dec_part = f"{number:.{decimal}f}".split(".")
    else:
        int_part = str(int(number))
        dec_part = ""

    if len(int_part) <= 3:
        formatted = int_part
    else:
        last3 = int_part[-3:]
        remaining = int_part[:-3]
        groups = []
        while remaining:
            groups.append(remaining[-2:])
            remaining = remaining[:-2]
        groups.reverse()
        formatted = ",".join(groups) + "," + last3

    result = f"{prefix}{formatted}"
    if dec_part:
        result += f".{dec_part}"
    if is_negative:
        result = f"-{result}"
    return result


def _find_excel_file() -> Path:
    """Locate the most recently modified Excel file in DATA_DIR.

    Returns:
        Path to the latest .xlsx file.

    Raises:
        FileNotFoundError: If no Excel file is found.
    """
    pattern = str(Path(DATA_DIR) / "*.xlsx")
    files = glob.glob(pattern)
    if not files:
        raise FileNotFoundError(f"No .xlsx files found in {DATA_DIR}/")
    return Path(max(files, key=os.path.getmtime))


# ---------------------------------------------------------------------------
# DATA LOADING
# ---------------------------------------------------------------------------


def load_data() -> pd.DataFrame:
    """Load and clean Excel data (no Streamlit dependency).

    Reads the latest .xlsx from DATA_DIR, renames columns to standard names,
    excludes self-counter dealers, and uses pre-calculated Points from Excel.

    Returns:
        Cleaned DataFrame with slab assignments.
    """
    file_path = _find_excel_file()
    log_info(f"Loading data from {file_path}")

    df = pd.read_excel(
        file_path,
        sheet_name=SHEET_NAME,
        header=HEADER_ROW,
        engine="openpyxl",
    )

    # --- Rename columns to standard names ---
    rename_map = {k: v for k, v in COLUMN_MAP.items() if k in df.columns}
    df = df.rename(columns=rename_map)
    log_info(f"Renamed columns: {list(rename_map.values())}")

    # --- Exclude self-counter dealers ---
    if "Self Counter" in df.columns:
        before = len(df)
        df["Self Counter"] = df["Self Counter"].fillna("").astype(str).str.strip().str.title()
        df = df[df["Self Counter"] != "Yes"].copy()
        excluded = before - len(df)
        log_info(f"Excluded {excluded} self-counter dealers ({len(df)} remaining)")

    # --- Coerce numeric columns ---
    numeric_cols = [
        "Qualified Volume", "Total Site Volume", "Qualified Points",
        "Shop Volume", "Site Volume", "Q4 Volume", "Unique Site >200 MT",
    ]
    for col in numeric_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0.0)

    # --- Clean text columns ---
    text_cols = ["Dealer Name", "Distributor Name", "State", "District", "Zone"]
    for col in text_cols:
        if col in df.columns:
            df[col] = (
                df[col]
                .fillna("")
                .astype(str)
                .str.strip()
                .str.title()
                .replace({"Nan": "", "None": "", "0": "", "0.0": ""})
            )

    # --- Total Volume (Shop + Site) ---
    if "Shop Volume" in df.columns and "Site Volume" in df.columns:
        df["Total Volume"] = df["Shop Volume"] + df["Site Volume"]
    elif "Q4 Volume" in df.columns:
        df["Total Volume"] = df["Q4 Volume"]
    else:
        df["Total Volume"] = 0.0

    # --- Derive Qualified Slab from pre-calculated Points ---
    if "Qualified Points" in df.columns:
        df["Qualified Slab"] = df["Qualified Points"].apply(assign_slab)
    elif "Current Slab" in df.columns:
        df["Current Slab"] = df["Current Slab"].fillna("-").astype(str).str.strip()
        df["Qualified Slab"] = df["Current Slab"].map(SLAB_CODE_MAP).fillna("Unqualified")
    else:
        df["Qualified Slab"] = "Unqualified"

    # --- Next slab & points gap ---
    df["Next Upgrade Slab"] = df["Qualified Slab"].apply(get_next_slab)
    if "Qualified Points" in df.columns:
        df["Points to Next Slab"] = df.apply(
            lambda row: points_to_next(row["Qualified Points"], row["Qualified Slab"]),
            axis=1,
        )

    log_info(f"Loaded {len(df)} rows")
    return df


# ---------------------------------------------------------------------------
# EXCEL REPORT GENERATION
# ---------------------------------------------------------------------------

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _auto_width(ws: object) -> None:
    """Auto-adjust column widths based on content.

    Args:
        ws: openpyxl worksheet object.
    """
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 40)


def _style_header_row(ws: object, num_cols: int) -> None:
    """Apply header styling to the first row.

    Args:
        ws: openpyxl worksheet object.
        num_cols: Number of columns to style.
    """
    for col_idx in range(1, num_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER


def generate_report(df: pd.DataFrame) -> Path:
    """Generate a formatted Excel slab report.

    Args:
        df: Cleaned DataFrame with slab assignments.

    Returns:
        Path to the generated report file.
    """
    output_dir = Path(OUTPUT_DIR)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "slab_report.xlsx"

    wb = Workbook()

    # --- Sheet 1: Slab Summary ---
    ws_summary = wb.active
    ws_summary.title = "Slab Summary"
    ws_summary.sheet_properties.tabColor = "1F4E79"

    headers = ["Slab", "Points Range", "Dealer Count", "Total Volume",
               "Qual. Shop Vol.", "Qual. Site Vol.", "Qualified Volume",
               "Total Points", "Gift"]
    ws_summary.append(headers)

    for cfg in SLAB_CONFIG:
        slab_df = df[df["Qualified Slab"] == cfg["slab"]] if "Qualified Slab" in df.columns else pd.DataFrame()
        count = len(slab_df)
        vol = slab_df["Total Volume"].sum() if "Total Volume" in slab_df.columns else 0
        shop_vol = slab_df["Shop Volume"].sum() if "Shop Volume" in slab_df.columns else 0
        site_vol = slab_df["Site Volume"].sum() if "Site Volume" in slab_df.columns else 0
        qual = slab_df["Qualified Volume"].sum() if "Qualified Volume" in slab_df.columns else 0
        pts = slab_df["Qualified Points"].sum() if "Qualified Points" in slab_df.columns else 0
        ws_summary.append([
            cfg["slab"], cfg["range"], count, int(round(vol)),
            int(round(shop_vol)), int(round(site_vol)),
            int(round(qual)), int(round(pts)), cfg["gift_full"],
        ])

    _style_header_row(ws_summary, len(headers))
    _auto_width(ws_summary)
    ws_summary.freeze_panes = "A2"

    # --- Sheet 2: Dealer Detail ---
    ws_detail = wb.create_sheet("Dealer Detail")
    ws_detail.sheet_properties.tabColor = "10B981"

    detail_cols = [
        c for c in ["Dealer Name", "Distributor Name", "State", "Zone",
                     "Shop Volume", "Site Volume", "Total Volume",
                     "Qualified Slab", "Next Upgrade Slab",
                     "Points to Next Slab", "Qualified Points"]
        if c in df.columns
    ]

    # Build display with renamed headers
    header_rename: dict[str, str] = {
        "Shop Volume": "Qual. Shop Vol.",
        "Site Volume": "Qual. Site Vol.",
        "Next Upgrade Slab": "Next Slab",
    }
    display_headers = [header_rename.get(c, c) for c in detail_cols]
    # Replace Points to Next Slab with Vol. to Achieve
    if "Points to Next Slab" in display_headers:
        display_headers[display_headers.index("Points to Next Slab")] = "Vol. to Achieve"
    ws_detail.append(display_headers)

    for _, row in df[detail_cols].iterrows():
        values = []
        for col, val in zip(detail_cols, row.values):
            if col == "Points to Next Slab":
                # Compute Vol. to Achieve from volume thresholds instead
                slab = row.get("Qualified Slab", "Unqualified")
                total_vol = float(row.get("Total Volume", 0) if pd.notna(row.get("Total Volume", 0)) else 0)
                next_vol = NEXT_SLAB_VOLUME.get(slab)
                vol_to_achieve = max(next_vol - total_vol, 0) if next_vol is not None else 0
                values.append(int(round(vol_to_achieve)))
            elif col in ("Shop Volume", "Site Volume", "Total Volume",
                         "Qualified Points"):
                values.append(int(round(float(val if pd.notna(val) else 0))))
            elif col == "Next Upgrade Slab":
                values.append(str(val) if pd.notna(val) else "-")
            else:
                values.append(val)
        ws_detail.append(values)

    _style_header_row(ws_detail, len(detail_cols))
    _auto_width(ws_detail)
    ws_detail.freeze_panes = "A2"

    # --- Sheet 3: Near-Upgrade ---
    ws_upgrade = wb.create_sheet("Near-Upgrade")
    ws_upgrade.sheet_properties.tabColor = "F59E0B"

    upgrade_headers = [
        "Dealer Name", "Distributor Name", "State", "Zone",
        "Qualified Points", "Qualified Slab", "Next Slab", "Pts to Upgrade",
        "Total Volume",
    ]
    ws_upgrade.append(upgrade_headers)

    # Filter to upgradable dealers (exclude top slab) within 500 pts
    last_slab = SLAB_CONFIG[-1]["slab"]
    upgradable = df[df["Qualified Slab"] != last_slab].copy()
    if "Qualified Points" in upgradable.columns:
        upgradable["Points Gap"] = upgradable.apply(
            lambda row: points_to_next(row["Qualified Points"], row["Qualified Slab"]) or 0,
            axis=1,
        )
        near = upgradable[upgradable["Points Gap"] <= 500].sort_values("Points Gap")

        for _, row in near.iterrows():
            ws_upgrade.append([
                row.get("Dealer Name", ""),
                row.get("Distributor Name", ""),
                row.get("State", ""),
                row.get("Zone", ""),
                int(round(float(row.get("Qualified Points", 0)))),
                row.get("Qualified Slab", ""),
                get_next_slab(row.get("Qualified Slab", "")) or "-",
                int(round(float(row.get("Points Gap", 0)))),
                int(round(float(row.get("Total Volume", 0)))),
            ])

    _style_header_row(ws_upgrade, len(upgrade_headers))
    _auto_width(ws_upgrade)
    ws_upgrade.freeze_panes = "A2"

    # --- Sheet 4: Distributor Performance ---
    ws_dist = wb.create_sheet("Distributor Performance")
    ws_dist.sheet_properties.tabColor = "6366F1"

    dist_headers = [
        "Distributor Name", "Dealers", "Total Volume",
        "Qual. Shop Vol.", "Qual. Site Vol.",
        "Avg Points", "Total Points", "Qual. Rate %",
    ] + SLAB_ORDER
    ws_dist.append(dist_headers)

    dist_agg = df.groupby("Distributor Name").agg(
        Dealers=("Dealer Name", "count"),
        Total_Volume=("Total Volume", "sum"),
        Avg_Points=("Qualified Points", "mean"),
        Total_Points=("Qualified Points", "sum"),
        Shop_Volume=("Shop Volume", "sum"),
        Site_Volume=("Site Volume", "sum"),
    )
    dist_agg = dist_agg[dist_agg.index.str.strip() != ""]
    dist_agg = dist_agg.sort_values("Dealers", ascending=False)

    slab_mix = df.groupby(["Distributor Name", "Qualified Slab"]).size().unstack(fill_value=0)
    for slab_name in SLAB_ORDER:
        if slab_name not in slab_mix.columns:
            slab_mix[slab_name] = 0
    slab_mix = slab_mix[SLAB_ORDER]

    if "Unqualified" in slab_mix.columns:
        slab_mix["Qualified Rate"] = (
            slab_mix.drop(columns=["Unqualified"]).sum(axis=1)
            / slab_mix.sum(axis=1) * 100
        )
    else:
        slab_mix["Qualified Rate"] = 100.0

    for dist_name in dist_agg.index:
        row_data = dist_agg.loc[dist_name]
        qual_rate = slab_mix.loc[dist_name, "Qualified Rate"] if dist_name in slab_mix.index else 0
        slab_counts = [
            int(slab_mix.loc[dist_name, s]) if dist_name in slab_mix.index else 0
            for s in SLAB_ORDER
        ]
        ws_dist.append([
            dist_name,
            int(row_data["Dealers"]),
            int(round(float(row_data["Total_Volume"]))),
            int(round(float(row_data["Shop_Volume"]))),
            int(round(float(row_data["Site_Volume"]))),
            int(round(float(row_data["Avg_Points"]))),
            int(round(float(row_data["Total_Points"]))),
            int(round(float(qual_rate))),
        ] + slab_counts)

    _style_header_row(ws_dist, len(dist_headers))
    _auto_width(ws_dist)
    ws_dist.freeze_panes = "A2"

    # --- Sheet 5: Zone Analysis ---
    ws_zone = wb.create_sheet("Zone Analysis")
    ws_zone.sheet_properties.tabColor = "10B981"

    zone_headers = [
        "Zone", "Dealers", "Total Volume",
        "Qual. Shop Vol.", "Qual. Site Vol.",
        "Avg Points", "Total Points", "Qual. Rate %",
    ] + SLAB_ORDER
    ws_zone.append(zone_headers)

    zone_agg = df.groupby("Zone").agg(
        Dealers=("Dealer Name", "count"),
        Total_Volume=("Total Volume", "sum"),
        Avg_Points=("Qualified Points", "mean"),
        Total_Points=("Qualified Points", "sum"),
        Shop_Volume=("Shop Volume", "sum"),
        Site_Volume=("Site Volume", "sum"),
    )
    zone_agg = zone_agg[zone_agg.index.str.strip() != ""]
    zone_agg = zone_agg.sort_values("Dealers", ascending=False)

    zone_slab_mix = (
        df[df["Zone"].str.strip() != ""]
        .groupby(["Zone", "Qualified Slab"])
        .size()
        .unstack(fill_value=0)
    )
    for slab_name in SLAB_ORDER:
        if slab_name not in zone_slab_mix.columns:
            zone_slab_mix[slab_name] = 0
    zone_slab_mix = zone_slab_mix[SLAB_ORDER]

    if "Unqualified" in zone_slab_mix.columns:
        zone_slab_mix["Qualified Rate"] = (
            zone_slab_mix.drop(columns=["Unqualified"]).sum(axis=1)
            / zone_slab_mix.sum(axis=1) * 100
        )
    else:
        zone_slab_mix["Qualified Rate"] = 100.0

    for zone_name in zone_agg.index:
        row_data = zone_agg.loc[zone_name]
        qual_rate = zone_slab_mix.loc[zone_name, "Qualified Rate"] if zone_name in zone_slab_mix.index else 0
        slab_counts = [
            int(zone_slab_mix.loc[zone_name, s]) if zone_name in zone_slab_mix.index else 0
            for s in SLAB_ORDER
        ]
        ws_zone.append([
            zone_name,
            int(row_data["Dealers"]),
            int(round(float(row_data["Total_Volume"]))),
            int(round(float(row_data["Shop_Volume"]))),
            int(round(float(row_data["Site_Volume"]))),
            int(round(float(row_data["Avg_Points"]))),
            int(round(float(row_data["Total_Points"]))),
            int(round(float(qual_rate))),
        ] + slab_counts)

    _style_header_row(ws_zone, len(zone_headers))
    _auto_width(ws_zone)
    ws_zone.freeze_panes = "A2"

    # Save
    wb.save(str(output_path))
    log_info(f"Report saved to {output_path}")
    return output_path


# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------


def main() -> None:
    """Load data and generate the slab report."""
    try:
        df = load_data()
        output = generate_report(df)
        print(f"\nReport generated: {output}")
    except FileNotFoundError as e:
        log_error(str(e))
        print(f"\nError: {e}")
        print("Please add an .xlsx file to the data/ directory.")
        sys.exit(1)


if __name__ == "__main__":
    main()
